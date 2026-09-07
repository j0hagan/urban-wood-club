#!/usr/bin/env python3
"""
Urban Wood Club - Monthly Research Data Export
================================================

Exports a snapshot of the tree database to a dated Excel workbook, for
tracking how many trees Delft has and how many are being felled over time.

USAGE (run from the api/ folder, once a month or whenever you want a fresh
snapshot):

    cd api
    python3 scripts/export_research_data.py

Each run creates a new folder at the project root:

    exports/<YYYY-MM-DD>/
        urban-wood-club-<YYYY-MM-DD>.xlsx   3 tabs: All Trees, Felling
                                             Permits, Findings
        summary.json                        machine-readable snapshot, used
                                             to compute month-over-month
                                             changes the next time you run
                                             this script

Requires openpyxl:  pip3 install openpyxl --break-system-packages

NOTE ON DATA SOURCE: this reads the LOCAL wrangler-dev D1 sqlite file
directly (the same one `wrangler dev` uses) rather than going through
`wrangler d1 execute`, because that command needs the Workers runtime
(workerd), which isn't guaranteed to be available everywhere this script
might run. Once the app is deployed for real and the live data lives in a
remote D1 database, swap find_db_path()/the sqlite3.connect() call below for
a `wrangler d1 execute --remote --json "SELECT * FROM trees"` style fetch
(via subprocess), keeping everything else the same.
"""

import sqlite3
import glob
import json
from pathlib import Path
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
API_DIR = SCRIPT_DIR.parent
PROJECT_ROOT = API_DIR.parent
EXPORTS_DIR = PROJECT_ROOT / "exports"

HEADER_FILL = PatternFill(start_color="2F5233", end_color="2F5233", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")

STATUS_EN = {
    "aangevraagd": "Requested",
    "verleend": "Granted",
    "geweigerd": "Refused",
}
STATUS_ORDER = [
    ("aangevraagd", "Requested"),
    ("verleend", "Granted"),
    ("geweigerd", "Refused"),
]


def find_db_path() -> Path:
    matches = glob.glob(str(API_DIR / ".wrangler/state/v3/d1/miniflare-D1DatabaseObject/*.sqlite"))
    if not matches:
        raise SystemExit(
            "Could not find the local D1 sqlite file under api/.wrangler/.\n"
            "Run `wrangler dev` at least once first (it creates this file on\n"
            "first start), then re-run this export."
        )
    return Path(matches[0])


def style_header_row(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize_columns(ws, max_width=60):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), max_width)


def export_trees(ws, cur):
    headers = ["ID", "Latitude", "Longitude", "Species (NL)", "Species (Latin)",
               "Monumental", "Planted Year", "Height Class", "Diameter (cm)",
               "Neighborhood", "Site Type", "Management Group", "Notes",
               "Source", "Source Ref", "Updated At"]
    ws.append(headers)
    cur.execute("""
        SELECT id, lat, lon, species_nl, species_lat, is_monumental, planted_year,
               height_class, diameter_cm, neighborhood, site_type, management_group,
               notes, source, source_ref, updated_at
        FROM trees
        ORDER BY neighborhood, species_nl
    """)
    n = 0
    for row in cur.fetchall():
        row = list(row)
        row[5] = "Yes" if row[5] else "No"
        ws.append(row)
        n += 1
    style_header_row(ws, len(headers))
    autosize_columns(ws)
    return n


def export_permits(ws, cur):
    headers = ["Publication ID", "Title (NL)", "Title (EN)", "Status (NL)", "Status (EN)",
               "Review Status", "Address", "Latitude", "Longitude", "Tree Count",
               "Species (reported)", "Reason for Felling", "Published At",
               "Source URL", "Created At"]
    ws.append(headers)
    cur.execute("""
        SELECT publication_id, title, title_en, status, review_status, address,
               lat, lon, tree_count, species, reason, published_at, source_url, created_at
        FROM felling_permits
        ORDER BY published_at DESC
    """)
    n = 0
    stats = {"aangevraagd": 0, "verleend": 0, "geweigerd": 0, "other": 0}
    trees_by_status = {"aangevraagd": 0, "verleend": 0, "geweigerd": 0, "other": 0}
    for row in cur.fetchall():
        (pub_id, title, title_en, status, review_status, address, lat, lon,
         tree_count, species, reason, published_at, source_url, created_at) = row
        status_en = STATUS_EN.get(status, status)
        ws.append([pub_id, title, title_en, status, status_en, review_status, address,
                   lat, lon, tree_count, species, reason, published_at, source_url, created_at])
        n += 1
        key = status if status in stats else "other"
        stats[key] += 1
        trees_by_status[key] += tree_count or 0
    style_header_row(ws, len(headers))
    autosize_columns(ws)
    return n, stats, trees_by_status


def gather_tree_stats(cur):
    cur.execute("SELECT count(*) FROM trees")
    total = cur.fetchone()[0]
    cur.execute("SELECT count(*) FROM trees WHERE is_monumental = 1")
    monumental = cur.fetchone()[0]
    cur.execute("""
        SELECT COALESCE(neighborhood, '(unknown)'), count(*)
        FROM trees GROUP BY neighborhood ORDER BY count(*) DESC
    """)
    by_neighborhood = cur.fetchall()
    cur.execute("""
        SELECT COALESCE(species_nl, '(unknown)'), count(*)
        FROM trees GROUP BY species_nl ORDER BY count(*) DESC LIMIT 15
    """)
    by_species = cur.fetchall()
    return {
        "total": total,
        "monumental": monumental,
        "by_neighborhood": by_neighborhood,
        "by_species_top15": by_species,
    }


def find_previous_summary(before_date_str):
    if not EXPORTS_DIR.exists():
        return None
    candidates = sorted(
        [p for p in EXPORTS_DIR.iterdir() if p.is_dir() and p.name < before_date_str],
        reverse=True,
    )
    for folder in candidates:
        summary_path = folder / "summary.json"
        if summary_path.exists():
            with open(summary_path) as f:
                return json.load(f)
    return None


def write_findings(ws, tree_stats, permit_count, permit_stats, trees_by_status, prev, run_date):
    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=12)

    def delta_str(curr, prev_val):
        if prev_val is None:
            return ""
        d = curr - prev_val
        if d == 0:
            return "no change since last export"
        sign = "+" if d > 0 else ""
        return f"{sign}{d} since last export"

    prev_tree_total = prev["tree_stats"]["total"] if prev else None
    prev_monumental = prev["tree_stats"]["monumental"] if prev else None
    prev_permit_count = prev["permit_count"] if prev else None
    prev_trees_by_status = prev.get("trees_by_status") if prev else None
    prev_run_date = prev["run_date"] if prev else None

    r = 1
    ws.cell(row=r, column=1, value="Urban Wood Club — Research Findings").font = title_font
    r += 1
    ws.cell(row=r, column=1, value=f"Export date: {run_date}")
    r += 1
    if prev_run_date:
        ws.cell(row=r, column=1, value=f"Compared against previous export: {prev_run_date}").font = Font(italic=True)
    else:
        ws.cell(row=r, column=1, value="This is the first export — no prior snapshot to compare against yet.").font = Font(italic=True)
    r += 2

    ws.cell(row=r, column=1, value="Tier 1 — Existing Trees").font = section_font
    r += 1
    ws.cell(row=r, column=1, value="Total trees tracked:")
    ws.cell(row=r, column=2, value=tree_stats["total"])
    ws.cell(row=r, column=3, value=delta_str(tree_stats["total"], prev_tree_total))
    r += 1
    ws.cell(row=r, column=1, value="Monumental trees:")
    ws.cell(row=r, column=2, value=tree_stats["monumental"])
    ws.cell(row=r, column=3, value=delta_str(tree_stats["monumental"], prev_monumental))
    r += 2

    ws.cell(row=r, column=1, value="Tier 2 — Felling Permits").font = section_font
    r += 1
    ws.cell(row=r, column=1, value="Total permits on record:")
    ws.cell(row=r, column=2, value=permit_count)
    ws.cell(row=r, column=3, value=delta_str(permit_count, prev_permit_count))
    r += 1
    for status, label in STATUS_ORDER:
        ws.cell(row=r, column=1, value=f"  {label} ({status}):")
        ws.cell(row=r, column=2, value=permit_stats.get(status, 0))
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Trees affected by permits (reported tree_count sums)").font = section_font
    r += 1
    for status, label in [("aangevraagd", "Pending decision"), ("verleend", "Granted for felling"), ("geweigerd", "Refused (protected)")]:
        prev_val = prev_trees_by_status.get(status) if prev_trees_by_status else None
        ws.cell(row=r, column=1, value=f"  {label}:")
        ws.cell(row=r, column=2, value=trees_by_status.get(status, 0))
        ws.cell(row=r, column=3, value=delta_str(trees_by_status.get(status, 0), prev_val))
        r += 1
    r += 2

    net = tree_stats["total"] - trees_by_status.get("verleend", 0)
    ws.cell(row=r, column=1, value="Est. trees remaining if all granted permits are executed:").font = bold
    ws.cell(row=r, column=2, value=net)
    r += 2

    ws.cell(row=r, column=1, value="Trees by neighborhood").font = section_font
    r += 1
    ws.cell(row=r, column=1, value="Neighborhood").font = bold
    ws.cell(row=r, column=2, value="Count").font = bold
    r += 1
    for name, count in tree_stats["by_neighborhood"]:
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=count)
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Top 15 species").font = section_font
    r += 1
    ws.cell(row=r, column=1, value="Species (NL)").font = bold
    ws.cell(row=r, column=2, value="Count").font = bold
    r += 1
    for name, count in tree_stats["by_species_top15"]:
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=count)
        r += 1

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 30


def main():
    run_date = datetime.now().strftime("%Y-%m-%d")
    out_dir = EXPORTS_DIR / run_date
    out_dir.mkdir(parents=True, exist_ok=True)

    db_path = find_db_path()
    con = sqlite3.connect(db_path)
    cur = con.cursor()

    wb = Workbook()
    ws_trees = wb.active
    ws_trees.title = "All Trees (Tier 1)"
    tree_count = export_trees(ws_trees, cur)

    ws_permits = wb.create_sheet("Felling Permits (Tier 2)")
    permit_count, permit_stats, trees_by_status = export_permits(ws_permits, cur)

    tree_stats = gather_tree_stats(cur)
    prev = find_previous_summary(run_date)

    ws_findings = wb.create_sheet("Findings")
    write_findings(ws_findings, tree_stats, permit_count, permit_stats, trees_by_status, prev, run_date)

    xlsx_path = out_dir / f"urban-wood-club-{run_date}.xlsx"
    wb.save(xlsx_path)

    summary = {
        "run_date": run_date,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "tree_stats": {"total": tree_stats["total"], "monumental": tree_stats["monumental"]},
        "permit_count": permit_count,
        "permit_stats": permit_stats,
        "trees_by_status": trees_by_status,
    }
    with open(out_dir / "summary.json", "w") as f:
        json.dump(summary, f, indent=2)

    con.close()

    print(f"Exported {tree_count} trees and {permit_count} felling permits.")
    print(f"Workbook: {xlsx_path}")
    if prev:
        print(f"Compared against previous export from {prev['run_date']}.")
    else:
        print("No previous export found -- this is the first snapshot (no comparison numbers yet).")


if __name__ == "__main__":
    main()
